from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import hashlib, json, re
from content import TEAM, TITLE, SUBTITLE, TEXT, ABSTRACTS

ROOT = Path(__file__).resolve().parents[1]
REF = Path(r'C:\Users\dhiraj\Downloads\IT DS CEP Documentation for gpt.docx')
OUT = ROOT / 'final'
OUT.mkdir(exist_ok=True)
source = Document(REF)
manifest = {'reference':str(REF),'sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'parts':{}}
with ZipFile(REF) as z:
    for n in z.namelist():
        manifest['parts'][n] = {'size':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()}
(ROOT/'work'/'manifest.json').write_text(json.dumps(manifest,indent=2))

CHAPTERS = [
('INTRODUCTION',['Project Background','Problem Statement','Need for the Project','Objectives','Scope','SDG Alignment']),
('REQUIREMENT ANALYSIS AND LITERATURE REVIEW',['Methodology Adopted','Stakeholder Interactions','Domain Overview','Existing Process Analysis','Existing Solutions and Related Work','Technologies Studied','Research Gap and Need for Proposed Solution','Challenges Identified']),
('PROPOSED SOLUTION',['Solution Overview','System Architecture','Process Flow Diagram','Module Description','Advantages of Proposed Solution']),
('SYSTEM DESIGN',['System Design Overview','Use Case Diagram','Workflow Diagram','Module Design']),
('PROJECT IMPLEMENTATION',['Development Methodology','Module-wise Implementation','User Interface Solution','Screenshots of Developed Solution']),
('TESTING AND VALIDATION',['Test Cases','Test Results','User Acceptance Testing','Feedback from Stakeholders','Performance Evaluation']),
('RESULTS AND IMPACT ASSESSMENT',['Deliverables Submitted','Results Achieved','Impact on Organization','SDG Contribution']),
('CONCLUSION AND FUTURE SCOPE',['Conclusion','Key Findings','Recommendations','Future Enhancements'])]

STACK = [
['React','Component-based frontend and role-specific pages'],
['TypeScript','Typed models and compile-time checks for frontend and server'],
['Vite','Frontend development server and production bundling'],
['Tailwind CSS','Interface layout and responsive styling'],
['Recharts','Research charts and visual summaries'],
['Node.js and Express','Runtime, API routes, authentication and authorization'],
['SQLite through node:sqlite','Persistent relational records, foreign keys and transactions'],
['tsx and oxlint','TypeScript execution and lint checks'],
['Google Forms','Manual circulation and collection of real student field responses']]
MODULES = [
['Authentication and access','Sessions, login/logout, role checks and account isolation'],
['Student profile and matching','Profile preferences, skills and explained opportunity relevance'],
['Opportunities and applications','Listing discovery, saved openings, submission and review status'],
['Workshops','Session management, capacity, enrolment and attendance'],
['Feedback and guidance','Student requests and coordinator responses'],
['Research','Restricted response records and public aggregate reporting'],
['Study and visit records','Three-visit documentation, action plans and evidence fields'],
['Administration','Protected content, approved diagnostics and backup/restore']]
ROLES = [
['Public visitor','Permitted public pages, opportunity information and aggregate research'],
['Student','Own profile, saved listings, applications, enrolments, feedback and guidance'],
['Placement Coordinator','Operational management, decisions, rosters and restricted research'],
['Administrator','Protected configuration, diagnostics, backups and authorized management']]
DB = [
['users; sessions','Account identities, roles and active authentication sessions'],
['student_profiles','Student-specific skills, interests and academic information'],
['opportunities; applications','Listings linked to student submissions and review decisions'],
['saved_opportunities','Student-specific saved listings'],
['training_sessions; workshop_enrollments','Workshop information and participation relationships'],
['feedback_submissions; guidance_requests','Requests and staff follow-up'],
['survey_responses','Structured responses with provenance and controlled access'],
['employer_outreach','Staff-managed employer contact records'],
['system_documents; app_metadata','Content documents and database initialisation state']]
WEIGHTS = [
['Department','40','Awarded when the listing includes the profile department'],
['Preferred role','25','Case-insensitive containment between a preferred role and listing title'],
['Skill overlap','Up to 20','Seven points per overlapping listed skill, capped at 20'],
['Location','10','Preferred location matches, or the listing uses remote work'],
['Academic cutoff','5','Entered percentage meets a detected two-digit percentage cutoff']]
CASES = [
['Authentication','Invalid credentials and unauthenticated requests','Rejected'],
['Role authorization','Student requests staff-only resources','Rejected'],
['Account isolation','One student requests another student’s records','Private records not exposed'],
['Application lifecycle','Submit, reject duplicate, record coordinator decision','Expected records and status retained'],
['Workshop lifecycle','Enrol, reject duplicate, cancel, enforce capacity','Participation rules enforced'],
['Research privacy','Public/student access to raw responses','Blocked; aggregate reporting available'],
['Content and restoration','Malformed documents and invalid relationships','Rejected; valid restoration supported'],
['Diagnostics','Client-supplied SQL and approved query identifiers','Raw SQL blocked; approved queries work'],
['Production lifecycle','Fresh database, restart and admin provisioning','Startup and provisioning checks passed']]
RESULTS = [['Acceptance','37','Passed'],['Workflow and security','97','Passed'],['Production lifecycle','14','Passed'],['Production serving','9','Passed'],['Total','157','Passed']]
DELIVER = [['Web portal','Student, coordinator and administrator interfaces'],['Backend and database','API, schema, authentication and persistent workflows'],['Verification scripts','Acceptance, workflow, lifecycle and serving checks'],['Configuration and instructions','Environment template, startup and backup scripts'],['Project report','Individual report and required supporting annexures']]
FILES = [['src/pages','User-facing pages'],['src/components','Reusable interface components'],['src/context/AppContext.tsx','Shared session and application state'],['src/services/apiService.ts','Client API communication'],['src/utils/matching.ts','Rule-based matching function'],['server/index.ts','API and frontend-serving entry point'],['server/db.ts','Schema and database initialisation'],['server/validation.ts','Input and restoration validators'],['server/config.ts','Environment and runtime configuration'],['scripts','Tests, provisioning and backup utilities'],['README.md; .env.example','Setup instructions and configuration examples']]
QUESTIONS = [
('Which department or course are you studying in?','Computer Science (CS); Information Technology (IT); Data Science (DS).'),
('What is your current year of study?','First; Second; Third; Fourth; Other.'),
('What job role or career domain are you mainly interested in?','Open text response.'),
('How many campus recruitment drives offered roles relevant to your preferred domain during the current academic year?','0; 1–2; 3 or more; Not sure.'),
('Is the placement cell currently meeting your placement-related needs?','Yes; Partly; No; Not enough experience to judge.'),
('How often do you receive placement information early enough to prepare and apply?','Always; Sometimes; Rarely; Never; Not applicable.'),
('How useful has the placement training you attended been for your preferred role?','1–5, where 1 is Not useful and 5 is Very useful; Have not attended.'),
('What is the biggest barrier you face in getting a suitable placement?','Limited relevant roles; Eligibility restrictions; Skill gaps; Late information; Interview preparation; Location or pay mismatch; Other; No major barrier.'),
('Which placement support do you need most urgently?','Relevant employer connections; Technical training; Aptitude preparation; Resume and interview help; Career guidance; Timely alerts; Internship support; Other.'),
('How satisfied are you with overall placement support?','1–5, where 1 is Very dissatisfied and 5 is Very satisfied; Not enough experience to judge.'),
('Briefly describe your experience with the placement process.','Open text response.'),
('What one change would most improve placement support for you?','Open text response.')]

def field(p, instruction, cache=''):
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); r._r.append(a)
    r=p.add_run(); a=OxmlElement('w:instrText'); a.set(qn('xml:space'),'preserve'); a.text=' '+instruction+' '; r._r.append(a)
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'separate'); r._r.append(a)
    if cache: p.add_run(cache)
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'end'); r._r.append(a)

for who,(name,roll) in enumerate(TEAM):
    d=Document(REF)
    # A working copy retains template styles, logos, relationships and document furniture.
    for el in list(d._element.body):
        if el.tag != qn('w:sectPr'): d._element.body.remove(el)
    for st in ['Title','Front Heading','Body Small','Caption','TOC Heading']:
        if st not in d.styles: d.styles.add_style(st,WD_STYLE_TYPE.PARAGRAPH)
    for st in d.styles:
        if st.type==WD_STYLE_TYPE.PARAGRAPH:
            st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
            st.font.size=Pt(11)
            if st._element.rPr is not None:
                for c in st._element.rPr.findall(qn('w:color')):
                    for k in [qn('w:themeColor'),qn('w:themeTint'),qn('w:themeShade')]: c.attrib.pop(k,None)
    normal=d.styles['Normal']; normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control=True
    normal.paragraph_format.keep_together=True
    for tn in ['toc 1','toc 2','toc 3','toc 4','table of figures']:
        if tn in d.styles:
            st=d.styles[tn];st.paragraph_format.line_spacing=1.5
            st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0)
    for hn,size in [('Heading 1',18),('Heading 2',18),('Heading 3',16)]:
        st=d.styles[hn]; st.font.size=Pt(size); st.font.bold=True
        pf=st.paragraph_format; pf.keep_with_next=True; pf.space_before=Pt(12); pf.space_after=Pt(6); pf.line_spacing=1.5
        pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
        np=st._element.find(qn('w:pPr'))
        if np is not None:
            for x in list(np):
                if x.tag in [qn('w:numPr'),qn('w:pBdr')]: np.remove(x)
    d.styles['Heading 1'].paragraph_format.page_break_before=True
    d.styles['Front Heading'].font.size=Pt(18); d.styles['Front Heading'].font.bold=True
    d.styles['Front Heading'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.styles['Front Heading'].paragraph_format.space_after=Pt(18)
    d.styles['Front Heading'].paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(10); d.styles['Caption'].font.italic=True
    d.styles['Caption'].paragraph_format.line_spacing=1.15
    d.styles['Caption'].paragraph_format.space_after=Pt(8)
    d.styles['Body Small'].font.size=Pt(10)
    for sec in d.sections:
        sec.page_width=Inches(8.268); sec.page_height=Inches(11.693)
        sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
        sec.header_distance=sec.footer_distance=Inches(.5)
        sec.different_first_page_header_footer=True
    def para(text='',style=None):
        return d.add_paragraph(text,style)
    def h(text,level=2): return d.add_heading(text,level)
    def front(text):
        d.add_page_break(); return para(text,'Front Heading')
    def cap(label,text):
        p=para('', 'Caption'); p.add_run(label+' '); field(p,'SEQ '+label+' \\* ARABIC'); p.add_run(': '+text); return p
    def table(title,headers,rows,widths=None):
        p=cap('Table',title); p.paragraph_format.keep_with_next=True
        t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
        if widths is None: widths=[1.9,4.35] if len(headers)==2 else [2.5,1.05,2.7]
        for c,w in zip(t.columns,widths): c.width=Inches(w)
        for j,txt in enumerate(headers): t.rows[0].cells[j].text=txt
        for row in rows:
            for j,txt in enumerate(row): t.add_row() if j==0 else None; t.rows[-1].cells[j].text=str(txt)
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr(); x=OxmlElement('w:cantSplit');pr.append(x)
            if i==0: x=OxmlElement('w:tblHeader');pr.append(x)
            for j,c in enumerate(row.cells):
                c.width=Inches(widths[j]); c.vertical_alignment=1
                tcpr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
                tcpr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');margins.append(x)
                tcpr.append(margins)
                if i==0:
                    x=OxmlElement('w:shd');x.set(qn('w:fill'),'E7E6E6');tcpr.append(x)
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:r.font.size=Pt(11);r.bold=(i==0)
        para()
        return t
    def placeholder(title,kind='diagram',height=58):
        p=para('(add picture)' if kind in ['photo','screenshot'] else '[To be added]')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(height/2);p.paragraph_format.space_after=Pt(height/2);p.paragraph_format.keep_with_next=True
        cap('Figure',title+(' (placeholder)' if kind!='photo' else ''))
    # Cover uses the source order, typography and institutional logo.
    cover = [deepcopy(p._p) for p in source.paragraphs[:22]]
    for el in cover:d._element.body.insert(len(d._element.body)-1,el)
    replacements={0:'COMMUNITY ENGAGEMENT PROJECT (CEP)',1:'SKILL BRIDGE',8:name,9:'Roll Number: '+roll,11:'[To be added: Guide name]',12:'[To be added: Guide designation]',5:'BACHELOR OF SCIENCE\n[To be added: Information Technology / Data Science]'}
    for i,txt in replacements.items():
        p=d.paragraphs[i]; props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear();r=p.add_run(txt)
        if props is not None:r._r.insert(0,props)
    d.paragraphs[1].style=d.styles['Title']
    r=d.paragraphs[1].add_run('\n'+SUBTITLE);r.font.size=Pt(13);r.bold=False
    for i,p in enumerate(d.paragraphs):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.15
        for r in p.runs:r.font.name='Times New Roman'
    d.paragraphs[0].paragraph_format.space_before=Pt(8);d.paragraphs[0].paragraph_format.space_after=Pt(10)
    # Certificate retains the source's prescribed institutional and signature arrangement.
    d.add_page_break()
    p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(ROOT/'work'/'image1.png'),width=Inches(3.1))
    for txt in ['VIDYALANKAR SCHOOL OF INFORMATION TECHNOLOGY','(Autonomous College affiliated to University of Mumbai)','MUMBAI–MAHARASHTRA–400037','DEPARTMENT OF COMPUTING']:
        p=para(txt);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
    para('CERTIFICATE','Front Heading')
    para(f'I hereby certify that {name} (Roll Number: {roll}), student of Vidyalankar School of Information Technology, studying in T.Y. B.Sc. [To be added: Information Technology / Data Science], Semester V, has completed a C.E. project titled “Skill Bridge — Community Employment and Skill Matching Portal” at Khalsa College, Matunga, during the academic year 2026–27.')
    sig=deepcopy(source.tables[0]._tbl);d._element.body.insert(len(d._element.body)-1,sig)
    front('ORGANIZATION VISIT OBSERVATION SHEET')
    para('Attach the Organization Visit Observation Sheet duly signed by the reporting officer and mentor.')
    para('Organization: Khalsa College, Matunga')
    para('Student: '+name+' | Roll Number: '+roll)
    para('[To be added: Signed observation sheet covering all three visits]')
    front('PROJECT GUIDE MEETING RECORD')
    para('Attach the Project Guide Meeting Record duly signed by the mentor.')
    para('Project: '+TITLE+' — '+SUBTITLE)
    para('Student: '+name+' | Roll Number: '+roll)
    para('[To be added: Signed guide meeting record]')
    front('DECLARATION')
    declarations=[
    f'I, {name}, a student of Vidyalankar School of Information Technology, T.Y. B.Sc. [To be added: Information Technology / Data Science], Semester V, declare that this report presents our Community Engagement Project, “Skill Bridge — Community Employment and Skill Matching Portal”, undertaken during the academic year 2026–27.',
    f'I, {name}, studying in T.Y. B.Sc. [To be added: Information Technology / Data Science], Semester V, at Vidyalankar School of Information Technology, submit this report on the Community Engagement Project “Skill Bridge — Community Employment and Skill Matching Portal” for the academic year 2026–27.',
    f'I, {name}, student of Vidyalankar School of Information Technology in T.Y. B.Sc. [To be added: Information Technology / Data Science], Semester V, declare that the work described in this report concerns our project “Skill Bridge — Community Employment and Skill Matching Portal”, completed as part of CEP in 2026–27.',
    f'I, {name}, of T.Y. B.Sc. [To be added: Information Technology / Data Science], Semester V, Vidyalankar School of Information Technology, present this report on “Skill Bridge — Community Employment and Skill Matching Portal” as the record of our CEP work for 2026–27.']
    para(declarations[who])
    para('The fieldwork consisted of three visits to Khalsa College, Matunga. Our team collected real student responses by manually circulating a Google Form. The technical descriptions in this report concern the shared application developed for the project.')
    para('Sources used in preparing the report are identified in the references. Supporting forms and signatures are to be attached in their prescribed locations.')
    p=para('Signature of the Student with Date: ____________________');p.paragraph_format.space_before=Pt(80)
    para(name+'\nRoll Number: '+roll)
    front('ACKNOWLEDGEMENT')
    acknowledgements=[
    ['I thank the University of Mumbai and Vidyalankar School of Information Technology for the opportunity to undertake this Community Engagement Project. It provided a setting in which classroom learning could be applied to a practical student-support problem.','I acknowledge our faculty mentor, [To be added: Guide name], and the support provided for the project. I also thank Khalsa College, Matunga, for the opportunity to carry out our three visits, and the students who took the time to respond to the Google Form.','I appreciate the shared effort of Vidhi Mahato, Gaurav Mhatre and Varun Patil throughout the team project. The report reflects a common body of fieldwork and software work, presented here in my own account.'],
    ['I am grateful to Vidyalankar School of Information Technology and the University of Mumbai for including community engagement in our academic work. The project allowed us to connect technical development with a student community.','My thanks go to our mentor, [To be added: Guide name], and to Khalsa College, Matunga, where the three visits took place. I also acknowledge the students who contributed their responses through the Google Form.','Working with Dhiraj Tendulkar, Gaurav Mhatre and Varun Patil made this a shared learning experience. I thank them for their part in the common project and recognise the contribution of everyone who supported its completion.'],
    ['I acknowledge the opportunity provided by the University of Mumbai and Vidyalankar School of Information Technology to undertake this CEP project. It helped bring together fieldwork, system design and technical verification.','I thank [To be added: Guide name], our faculty mentor, and Khalsa College, Matunga, for their support of the project process. The students who completed the circulated Google Form made an important contribution to the fieldwork.','I am thankful to my teammates, Dhiraj Tendulkar, Vidhi Mahato and Varun Patil. Our shared work forms the basis of this report, and I value the cooperation involved in developing and reviewing the application.'],
    ['I thank Vidyalankar School of Information Technology and the University of Mumbai for the opportunity to complete a project with a community setting. The work gave practical meaning to the development and database concepts studied in our course.','I acknowledge our faculty mentor, [To be added: Guide name], and thank Khalsa College, Matunga, for hosting the project visits. I am also grateful to the students who shared information through the Google Form.','My teammates, Dhiraj Tendulkar, Vidhi Mahato and Gaurav Mhatre, contributed to the common project effort. I thank them and the faculty members who supported our work.']]
    for t in acknowledgements[who]:para(t)
    front('ABSTRACT')
    for t in ABSTRACTS[who].split('\n\n'):para(t)
    front('TABLE OF CONTENT')
    field(para(),'TOC \\o "1-2" \\h \\z \\u','[Update table of contents in Word]')
    front('TABLE OF FIGURES')
    field(para(),'TOC \\h \\z \\c "Figure"','[Update table of figures in Word]')
    front('TABLE OF TABLES')
    field(para(),'TOC \\h \\z \\c "Table"','[Update table of tables in Word]')
    sec=d.add_section(WD_SECTION_START.NEW_PAGE);sec.different_first_page_header_footer=False
    sec.footer.is_linked_to_previous=False
    for j,s in enumerate(d.sections):
        for el in list(s.footer._element):s.footer._element.remove(el)
        p=s.footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(p,'PAGE')
        pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'lowerRoman' if j==0 else 'decimal');pg.set(qn('w:start'),'1');s._sectPr.append(pg)

    for ci,(chapter,sections) in enumerate(CHAPTERS,1):
        p=h(f'CHAPTER {ci}: {chapter}',1)
        if ci==1:p.paragraph_format.page_break_before=False
        for si,heading in enumerate(sections,1):
            key=f'{ci}.{si}';hp=h(key+' '+heading)
            if key=='5.4':hp.paragraph_format.page_break_before=True
            compact_dhiraj={
            '6.2':'The latest run passed all four automated suites, the TypeScript check and the production build. Lint completed with warnings. These results apply to the tested behaviour; browser interaction and live-host verification remain separate.',
            '6.3':'Signed UAT records, participant details and observed outcomes are [To be added]. Actual users should verify the intended tasks in the browser. Automated acceptance checks do not establish formal acceptance by the college.',
            '6.4':'The real Google Form response summary and approved participant quotations are [To be added]. Annexure D contains a blank reporting-officer form. No unsupported statement is attributed to a student or staff member.',
            '6.5':'The tests establish functional behaviour, not measured speed or scale. Response times, throughput and concurrent-user capacity are [To be added]. A documented workload and target-host measurements are required before making performance claims.'}
            para(compact_dhiraj.get(key,TEXT[key][who]) if who==0 else TEXT[key][who])
            if key=='1.1':table('Project team',['Member','Roll number'],TEAM,[3.7,2.55])
            if key=='1.4':
                h('Primary Objectives',3)
                for t in ['Provide a distinct student workspace and coordinator workspace with server-enforced access.','Support discovery, saving and application for opportunities with visible status updates.','Connect workshop participation, feedback and guidance with coordinator follow-up.']:para('• '+t)
                h('Secondary Objectives',3)
                for t in ['Explain opportunity matching through understandable rules and reasons.','Preserve related records and reject invalid or unauthorized changes.','Provide aggregate research reporting while protecting individual responses.','Maintain repeatable tests and practical operating instructions.']:para('• '+t)
            if key=='2.1':table('Fieldwork record',['Item','Confirmed information or record'],[['Organization','Khalsa College, Matunga'],['Visits','Three visits, all at the same college'],['Collection method','Google Form manually circulated among students'],['Visit dates','[To be added]'],['Confirmed response count','[To be added]'],['Google Form link and response export','[To be added]'],['Verified findings and response summary','[To be added]']])
            if key=='2.6':table('Technologies used',['Technology','Role in Skill Bridge'],STACK)
            if key=='3.2':placeholder('System Architecture of Skill Bridge')
            if key=='3.3':placeholder('Process Flow of Skill Bridge')
            if key=='3.4':table('Functional modules',['Module','Responsibility'],MODULES)
            if key=='4.2':
                table('Actors and access',['Actor','Permitted scope'],ROLES)
                placeholder('Use Case Diagram of Skill Bridge')
            if key=='4.3':placeholder('Application and Workshop Workflow of Skill Bridge')
            if key=='4.4':
                h('Database records and relationships',3);table('Database entities',['Entity or group','Purpose'],DB)
                placeholder('Key Components of the Skill Bridge Database')
                h('Opportunity matching process',3)
                table('Matching score components',['Component','Points','Implemented rule'],WEIGHTS,[1.35,.8,4.1])
                para('The score is the sum of the awarded components, bounded between 0 and 100. Skill points are calculated as min(20, 7 × number of matching listing skills). A score of 75 or more is labelled High Fit; 50–74 is Moderate Fit; 30–49 is Partial Fit; lower scores are General Opportunity.')
                para('Text comparisons use case-insensitive containment, so the result depends on the terms entered in the profile and listing. The eligibility check detects a two-digit percentage cutoff and returns an eligibility caution where appropriate. It is not a complete eligibility parser. Students must review the original criteria even when the fit score is high.')
            if key=='5.2':
                table('Implementation mapping',['Source area','Implementation role'],FILES)
                h('Authentication and data integrity',3)
                para('Passwords are hashed with scrypt and a salt. Session tokens are carried in an HttpOnly cookie, and protected routes verify the session and role. Public registration is restricted to the student role. Database uniqueness constraints help prevent duplicate applications and enrolments. Restore operations validate supported content and relationships before transactional updates.')
                h('Research summaries',3)
                para('The reporting process groups responses by the requested filters and computes counts, distributions and rating summaries. Non-rated responses such as “Have not attended” are not treated as numeric satisfaction scores. Public reporting uses aggregate endpoints; individual responses are limited to authorized staff. The actual field-response totals and numerical findings are [To be added].')
            if key=='5.4':
                placeholder('Student Workspace of Skill Bridge','screenshot',80)
                placeholder('Placement Coordinator Workspace of Skill Bridge','screenshot',80)
                placeholder('Research Summary View of Skill Bridge','screenshot',80)
            if key=='6.1':table('Representative test cases',['Area','Scenario','Expected result'],CASES,[1.3,2.65,2.3])
            if key=='6.2':table('Recorded automated results',['Suite','Checks','Result'],RESULTS,[3.6,1.1,1.55])

    h('REFERENCES',1)
    h('Research Papers');para('[To be added: Verified research papers actually consulted, formatted in MLA style.]')
    h('Books Referred');para('[To be added: Books actually consulted, with verified MLA bibliographic details.]')
    h('Website');para('[To be added: Websites actually consulted and the topics referred to, with verified URLs and access details.]')
    para('Project source material:')
    for t in ['Skill Bridge Project Team. Skill Bridge: Community Employment and Skill Matching Portal. Project source code, 2026.','Skill Bridge Project Team. “Project Setup and Verification.” README.md, 2026.','Skill Bridge Project Team. “Opportunity Matching.” src/utils/matching.ts, 2026.','Skill Bridge Project Team. “Survey Questionnaire.” src/data/defaultSurveySeed.ts, 2026.','Skill Bridge Project Team. “Verification Suites.” scripts/verify-acceptance.ts, scripts/test-workflows.ts, scripts/test-production-startup.ts, and scripts/verify-production-serving.ts, 2026.']:
        p=para(t);p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)

    h('ANNEXURES',1)
    para('The annexures retain the labels prescribed in the supplied format: A, B, C, D, G and H.')
    h('Annexure A: Field Visit Photographs')
    for day in range(1,4):
        if day>1:d.add_page_break()
        h('Visit '+str(day)+' — Khalsa College, Matunga',3)
        para('Date of visit: [To be added]')
        placeholder('Visit '+str(day)+' at Khalsa College, Matunga','photo',280)
    h('ANNEXURE B: SURVEY QUESTIONNAIRE',1)
    para(['We collected real student inputs by manually circulating a Google Form. The questions below reproduce the questionnaire present in the project source. The original circulated form link and any differences from this version are [To be added].','Our field collection used a Google Form shared manually among students. This annexure records the questionnaire implemented in the project. The original form URL and confirmation of its final wording are [To be added].','The team used a manually circulated Google Form to gather genuine responses. The following questions are available in the project questionnaire. The original collection form and any version changes are [To be added].','Real student responses were collected through the Google Form distributed by the team. This annexure presents the project’s stored questionnaire. The circulated form link and final version confirmation are [To be added].'][who])
    for i,(q,opts) in enumerate(QUESTIONS,1):
        if i==7:d.add_page_break()
        p=para(f'{i}. {q}');p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
        para('Response options: '+opts)
        if i in [3,11,12]:para('________________________________________________________________\n________________________________________________________________')
    h('ANNEXURE C: INTERVIEW QUESTIONS',1)
    para('The project includes an interview guide for placement-coordination discussions. The questions below reproduce that guide. Interview dates, respondent identity, consent and recorded answers are [To be added]. Illustrative answers in the software are not reproduced as actual staff statements.')
    study=(ROOT.parent/'src/data/defaultStudy.ts').read_text(encoding='utf-8')
    guide=study.split('interviewGuide: [',1)[1].split('gapAnalysis:',1)[0]
    interview=re.findall(r"question:\s*'((?:\\.|[^'])*)'",guide)
    for i,q in enumerate(interview,1):para(f'{i}. '+q.replace("\\'","'"))
    h('ANNEXURE D: REPORTING OFFICER FEEDBACK FORM',1)
    para('Blank form for assessment of the prepared project')
    for label in ['Project: Skill Bridge — Community Employment and Skill Matching Portal','Organization: Khalsa College, Matunga','Student: '+name+' | Roll Number: '+roll,'Reporting officer name: ______________________________','Designation: _______________________________________','Date of review: _____________________________________']:
        para(label)
    for label in ['Relevance of the project to student/community needs','Usefulness and clarity of the developed solution','Observations on the demonstration','Suggested improvements','Overall assessment and acceptance status']:
        p=para(label);p.runs[0].bold=True;para('________________________________________________________________\n________________________________________________________________')
    para('Reporting officer signature: __________________________\nOrganization stamp, if required: ______________________')
    h('ANNEXURE G: SOURCE CODE REPOSITORY',1)
    para('Repository link: [To be added]\nSubmitted version or commit identifier: [To be added]\nContents uploaded and verified by: [To be added]')
    para('The following inventory identifies the project contents to include in the submission. It does not confirm that a repository has already been published.')
    table('Repository contents',['Source area','Contents'],FILES)
    para('Exclude passwords, private environment files, real student records and local database backups from any public repository. Provide authorised access to sensitive evidence separately.')
    h('ANNEXURE H: USER MANUAL',1)
    introductions=[
    'This manual follows the principal tasks of each role. Use the account provided for the intended role and keep credentials private. The actual hosting URL is [To be added].',
    'The instructions below describe how to use the completed workflows. Access depends on the signed-in account. The deployment address and support contact are [To be added].',
    'Use this guide to follow the main student and staff activities. A user must sign in with the appropriate account before performing protected actions. The live website address is [To be added].',
    'This guide summarises routine operation of Skill Bridge. The available controls depend on the account role. The final portal URL and operational contact are [To be added].']
    para(introductions[who])
    manual=[
    ('Student workflow',[
    ['Sign in and review your profile, including department, skills and preferred roles.','Open opportunities, apply relevant filters and read the listing details.','Review matching reasons and the original eligibility requirements before applying.','Save an opening for later or submit one application; check its status in your workspace.','Use the workshop section to enrol, review participation or cancel where permitted.','Submit feedback or a guidance request and return to view the response. Sign out when finished.'],
    ['Start by signing in to the student account and checking the accuracy of profile details.','Browse the opportunities and use filters to narrow the list.','Read the fit explanation as an aid, then verify the stated eligibility and deadline.','Apply to a selected opening or save it; return to the application record for updates.','Enrol in an appropriate workshop if capacity is available.','Use feedback or guidance for additional support, and log out after the session.'],
    ['Enter the student workspace and complete the profile fields used for matching.','Find a listing through browsing or filtering and inspect its requirements.','Check the match reasons without treating the score as a selection guarantee.','Submit an application once and follow its recorded status.','Review workshop details before enrolling or cancelling participation.','Record a feedback or guidance request when needed, then sign out securely.'],
    ['Authenticate as a student and confirm that profile information is current.','Explore relevant opportunities and open their full details.','Compare the explanation with the original eligibility conditions.','Save or apply to the opportunity and monitor the stored application status.','Manage workshop participation through the available enrolment controls.','Use the support forms for feedback or guidance and end the session by logging out.']][who]),
    ('Coordinator workflow',[
    ['Sign in with a placement coordinator account.','Maintain opportunities and workshop records through the management views.','Review student applications and record decisions or feedback.','Use workshop rosters to maintain the available participation and attendance information.','Respond to feedback and guidance requests; review research records only for authorised work.','Confirm save messages and revisit the record if a request fails.'],
    ['Open the coordinator workspace using the assigned account.','Create or update opportunity and training information after checking its accuracy.','Work through the application queue and save the appropriate decision.','Review enrolments and attendance from the workshop roster.','Handle guidance and feedback requests through their management screens.','Check the returned status after each update and sign out when the work is complete.'],
    ['Use coordinator credentials to access the operational dashboard.','Keep listings and workshops current, including dates and capacity.','Review the application records and provide the relevant decision.','Maintain workshop participation and attendance from the available roster.','Respond to student support requests and use restricted research views responsibly.','Treat an error message as an unsuccessful action until the record is verified.'],
    ['Log in to the placement coordinator account.','Manage the published opportunities and workshop details.','Inspect submitted applications and store decisions with any relevant feedback.','Use the enrolment roster for participation and attendance work.','Follow up on guidance and feedback queues.','Verify successful updates and close the session through logout.']][who]),
    ('Administrator workflow',[
    ['Use an explicitly provisioned administrator account.','Maintain protected content and run only the approved diagnostic queries.','Export a backup before a sensitive maintenance operation.','Restore only a validated backup and check the resulting records.','Keep database backups and credentials in protected storage.'],
    ['Sign in through the provisioned administrative account.','Review configuration and use the registered diagnostics for maintenance.','Create a backup before performing sensitive changes.','Use the restore control with an appropriate backup file and confirm the result.','Protect the backup files and avoid sharing administrative credentials.'],
    ['Authenticate with the administrator account created for the deployment.','Use protected content and diagnostic tools only for necessary maintenance.','Take a backup before operations that could affect stored data.','Check a restoration result and the linked operational records.','Store backups securely and log out after administration.'],
    ['Access the administrative workspace using the assigned credentials.','Maintain system content and use the approved diagnostics.','Back up the database before sensitive maintenance.','Submit a suitable backup for restoration and inspect the returned outcome.','Restrict access to backup files and administrator credentials.']][who])]
    for title,steps in manual:
        h(title,3)
        for i,t in enumerate(steps,1):para(str(i)+'. '+t)
    h('Local setup and production preparation',3)
    for t in ['Install a Node.js version compatible with the project dependencies and native SQLite support, then run npm install in the project folder.','Use npm run dev for local development. The configured API and frontend addresses should be read from the terminal output.','Run npm run typecheck, npm run build and npm test before preparing a release.','For production, set NODE_ENV=production, a persistent DB_PATH and CLIENT_URL matching the actual website origin. Configure HTTPS and the intended reverse proxy.','Provision an administrator with explicit credentials using npm run admin:provision -- <username> <password>. Use protected credential handling in the deployment environment.','Run npm start after the build. Verify login, authorised changes and direct route refreshes on the actual host.','Use npm run backup:db for the database backup utility. Test recovery with a disposable copy before relying on a backup. Do not run demo reset commands against the real database.']:
        para(t)
    h('If an operation fails',3)
    para(['Read the displayed error and check the current record before repeating a submission. Confirm that the server is running and the account has the required role. For repeated failures, record the action and error for the administrator without sharing private student data.','Check whether the action was saved before trying again. Verify the account role and server connection, then use the reported message to guide the next step. Escalate persistent errors with a clear description and no confidential records.','An error should be investigated rather than treated as a successful save. Recheck the record, connection and account permissions. Report reproducible failures to the administrator with the relevant steps, keeping personal information private.','When a request fails, inspect the current state first to avoid duplicate work. Check connectivity and permissions, and report continuing problems with the exact action and message. Do not include private records in a public error report.'][who])

    # Normalise page furniture and fields; preserve source package resources.
    update=d.settings.element.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');d.settings.element.append(update)
    update.set(qn('w:val'),'true')
    for p in d.paragraphs:
        p.paragraph_format.widow_control=True
        if '\n' in p.text and p.style.name=='Normal':p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
    d.core_properties.author=name;d.core_properties.title=TITLE+' — CEP Report — '+name
    d.core_properties.subject=SUBTITLE;d.core_properties.comments=''
    filename=name.replace(' ','_')+'_CEP_Report.docx'
    d.save(OUT/filename)
    print(filename,'abstract words',len(ABSTRACTS[who].split()),'body words',sum(len(p.text.split()) for p in d.paragraphs))
