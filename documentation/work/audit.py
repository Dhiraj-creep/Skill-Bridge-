from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree
import hashlib,json,re
from content import TEAM
root=Path('documentation')
expected=[6,8,5,4,4,5,4,4]
for name,roll in TEAM:
 f=root/'final'/(name.replace(' ','_')+'_CEP_Report.docx');d=Document(f)
 headings=[p.text for p in d.paragraphs if p.style.name in ['Heading 1','Heading 2']]
 text='\n'.join(p.text for p in d.paragraphs)
 assert name in d.paragraphs[8].text and roll in d.paragraphs[9].text
 for c,count in enumerate(expected,1):
  assert any(h.startswith(f'CHAPTER {c}:') for h in headings)
  for s in range(1,count+1):assert any(h.startswith(f'{c}.{s} ') for h in headings)
 for a in ['B','C','D','G','H']:assert any(h.startswith('ANNEXURE '+a+':') for h in headings)
 assert text.count('(add picture)')==6
 assert 'Error! Reference' not in text and '\ufffd' not in text
 assert 'Khalsa College' in text and 'Google Form' in text
 assert len(d.sections)==2
 assert '157' in '\n'.join(cell.text for t in d.tables for row in t.rows for cell in row.cells)
 print(f.name,'PASS',len(headings),'headings',f.stat().st_size,'bytes')
ref=Path(r'C:\Users\dhiraj\Downloads\IT DS CEP Documentation for gpt.docx')
assert hashlib.sha256(ref.read_bytes()).hexdigest()==json.loads((root/'work/manifest.json').read_text())['sha256']
print('Source template unchanged. All four reports verified.')
